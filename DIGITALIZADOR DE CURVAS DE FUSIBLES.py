# -*- coding: utf-8 -*-
"""
Digitalizador de Curvas de Fusibles para ETAP - V5
--------------------------------------------------------
Herramienta independiente en Python para digitalizar curvas TCC/fusibles a partir
 de imágenes, definiendo ejes, límites y puntos seleccionados manualmente.

Cambios V4:
- Se agrega la sección 5. Ayuda rápida al lado de 1. Imagen.
- La sección Acciones pasa a ser 6. Acciones y queda al lado derecho de 2. Información del fusible.

Cambios V3:
- Se reubica la sección Acciones en una columna lateral junto a 1. Imagen para que el botón Generar Excel quede visible.

Cambios V2:
- Dos pestañas de captura: Mínima de Fusión y Total de Aclaramiento.
- Cada pestaña tiene su propia lista de puntos.
- Exportación a Excel con formato similar al archivo ejemplo:
    Fila 2: Mínima de Fusión / Total de Aclaramiento
    Fila 3: Nombre del fusible / curva
    Fila 4: Time / Current / Time / Current
    Columnas A:B para Mínima de Fusión y C:D para Total de Aclaramiento.
- Exportación opcional de una hoja de configuración y trazabilidad.

Uso típico para curvas de fusibles:
- Eje X: Corriente [A], normalmente logarítmico.
- Eje Y: Tiempo [s], normalmente logarítmico.
- En el Excel de salida se exporta como: Time, Current.

Requisitos:
    pip install pillow pandas openpyxl

Autor: Herramienta generada con ChatGPT para flujos técnicos ETAP / estudios eléctricos.
"""

import os
import math
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from datetime import datetime

import pandas as pd
from PIL import Image, ImageTk


APP_TITLE = "Digitalizador de Curvas de Fusibles para ETAP"
DEFAULT_X_LABEL = "Corriente"
DEFAULT_X_UNIT = "A"
DEFAULT_Y_LABEL = "Tiempo"
DEFAULT_Y_UNIT = "s"

CURVE_MIN_FUSION = "Mínima de Fusión"
CURVE_TOTAL_CLEARING = "Total de Aclaramiento"
CURVE_KEYS = [CURVE_MIN_FUSION, CURVE_TOTAL_CLEARING]
CURVE_COLORS = {
    CURVE_MIN_FUSION: {"line": "#0B63F6", "fill": "#DCEBFF"},
    CURVE_TOTAL_CLEARING: {"line": "#00A676", "fill": "#DDF7EC"},
}

# Paleta visual de la interfaz.
COLOR_APP_BG = "#F4F7FB"
COLOR_PANEL_BG = "#FFFFFF"
COLOR_HEADER = "#0A2342"
COLOR_HEADER_2 = "#102E56"
COLOR_TEXT = "#102033"
COLOR_MUTED = "#5F6B7A"
COLOR_BORDER = "#D8E0EA"
COLOR_PRIMARY = "#0B63F6"
COLOR_PRIMARY_DARK = "#084EC1"
COLOR_SUCCESS = "#129A51"
COLOR_DANGER = "#D92D20"


# =============================================================================
# UTILIDADES NUMERICAS
# =============================================================================

def parse_float(text, field_name="valor"):
    try:
        value = str(text).replace(",", ".").strip()
        if value == "":
            raise ValueError
        return float(value)
    except Exception:
        raise ValueError("Ingrese un número válido para " + field_name + ".")


def validate_positive(value, field_name):
    if value <= 0:
        raise ValueError(field_name + " debe ser mayor que cero para escala logarítmica.")


def interp_linear(pixel, p1, p2, v1, v2):
    if abs(p2 - p1) < 1e-9:
        raise ValueError("Los puntos de calibración no pueden tener la misma coordenada de píxel.")
    ratio = (pixel - p1) / (p2 - p1)
    return v1 + ratio * (v2 - v1)


def interp_log(pixel, p1, p2, v1, v2):
    validate_positive(v1, "Valor mínimo")
    validate_positive(v2, "Valor máximo")
    log_v1 = math.log10(v1)
    log_v2 = math.log10(v2)
    log_value = interp_linear(pixel, p1, p2, log_v1, log_v2)
    return 10 ** log_value


def pixel_to_value(pixel, p1, p2, v1, v2, scale_type):
    if scale_type.lower() == "logarítmica":
        return interp_log(pixel, p1, p2, v1, v2)
    return interp_linear(pixel, p1, p2, v1, v2)


def format_number(value):
    try:
        value = float(value)
    except Exception:
        return str(value)
    if value == 0:
        return "0"
    abs_v = abs(value)
    if abs_v >= 1000 or abs_v < 0.01:
        return f"{value:.6g}"
    return f"{value:.6f}".rstrip("0").rstrip(".")


def safe_sheet_name(text):
    name = str(text).strip() or "Fuse 1"
    for ch in ['\\', '/', ':', '*', '?', '"', '<', '>', '|', '[', ']']:
        name = name.replace(ch, "_")
    return name[:31]


# =============================================================================
# CLASE PRINCIPAL
# =============================================================================

class FuseCurveDigitizerApp:
    def __init__(self, root):
        self.root = root
        self.root.title(APP_TITLE)
        self.root.geometry("1460x880")
        self.root.minsize(1220, 760)

        self.image_path = None
        self.original_image = None
        self.tk_image = None
        self.image_on_canvas = None
        self.image_scale = 1.0
        self.image_offset_x = 0
        self.image_offset_y = 0

        # Calibración en píxeles reales de la imagen original.
        self.calibration = {
            "x_min_pixel": None,
            "x_max_pixel": None,
            "y_min_pixel": None,
            "y_max_pixel": None,
            "x_min_value": None,
            "x_max_value": None,
            "y_min_value": None,
            "y_max_value": None,
        }

        self.mode = tk.StringVar(value="seleccionar_puntos")
        self.x_scale_type = tk.StringVar(value="Logarítmica")
        self.y_scale_type = tk.StringVar(value="Logarítmica")
        self.fuse_name = tk.StringVar(value="5E")
        self.sheet_name = tk.StringVar(value="Fuse 1")
        self.x_label = tk.StringVar(value=DEFAULT_X_LABEL)
        self.x_unit = tk.StringVar(value=DEFAULT_X_UNIT)
        self.y_label = tk.StringVar(value=DEFAULT_Y_LABEL)
        self.y_unit = tk.StringVar(value=DEFAULT_Y_UNIT)
        self.include_config_sheet = tk.BooleanVar(value=True)

        self.x_min_entry = tk.StringVar(value="10")
        self.x_max_entry = tk.StringVar(value="10000")
        self.y_min_entry = tk.StringVar(value="0.01")
        self.y_max_entry = tk.StringVar(value="1000")

        self.active_curve = tk.StringVar(value=CURVE_MIN_FUSION)
        self.points = {curve: [] for curve in CURVE_KEYS}
        self.point_counter = {curve: 1 for curve in CURVE_KEYS}
        self.marker_ids = []
        self.calibration_marker_ids = []
        self.trees = {}

        self._build_ui()
        self._bind_shortcuts()

    # -------------------------------------------------------------------------
    # INTERFAZ
    # -------------------------------------------------------------------------
    def _build_ui(self):
        self._configure_style()
        self.root.configure(bg=COLOR_APP_BG)

        header = tk.Frame(self.root, bg=COLOR_HEADER, height=58)
        header.pack(fill="x", side="top")
        header.pack_propagate(False)

        icon_box = tk.Label(
            header,
            text="⎍",
            bg=COLOR_PRIMARY,
            fg="white",
            font=("Segoe UI", 16, "bold"),
            width=3
        )
        icon_box.pack(side="left", padx=(14, 10), pady=10)

        title_block = tk.Frame(header, bg=COLOR_HEADER)
        title_block.pack(side="left", fill="y")
        tk.Label(
            title_block,
            text=APP_TITLE,
            bg=COLOR_HEADER,
            fg="white",
            font=("Segoe UI", 15, "bold")
        ).pack(anchor="w", pady=(8, 0))
        tk.Label(
            title_block,
            text="Digitalización manual de curvas TCC · Mínima de Fusión · Total de Aclaramiento",
            bg=COLOR_HEADER,
            fg="#C9D7EA",
            font=("Segoe UI", 8)
        ).pack(anchor="w", pady=(0, 6))

        main = ttk.Frame(self.root, style="App.TFrame")
        main.pack(fill="both", expand=True, padx=10, pady=10)

        self.left_panel = ttk.Frame(main, width=540, style="App.TFrame")
        self.left_panel.pack(side="left", fill="y", padx=(0, 8))
        self.left_panel.pack_propagate(False)

        right_panel = ttk.Frame(main, style="App.TFrame")
        right_panel.pack(side="right", fill="both", expand=True, padx=(0, 0))

        self._build_left_panel(self.left_panel)
        self._build_canvas_area(right_panel)
        self._build_tabs_area(right_panel)

    def _configure_style(self):
        style = ttk.Style()
        try:
            style.theme_use("clam")
        except Exception:
            pass

        style.configure("App.TFrame", background=COLOR_APP_BG)
        style.configure("Panel.TFrame", background=COLOR_PANEL_BG)
        style.configure("TLabel", background=COLOR_PANEL_BG, foreground=COLOR_TEXT, font=("Segoe UI", 9))
        style.configure("Muted.TLabel", background=COLOR_PANEL_BG, foreground=COLOR_MUTED, font=("Segoe UI", 8))
        style.configure("Title.TLabel", background=COLOR_PANEL_BG, foreground=COLOR_TEXT, font=("Segoe UI", 14, "bold"))
        style.configure("Section.TLabelframe", background=COLOR_PANEL_BG, bordercolor=COLOR_BORDER, relief="solid")
        style.configure("Section.TLabelframe.Label", background=COLOR_APP_BG, foreground=COLOR_TEXT, font=("Segoe UI", 10, "bold"))
        style.configure("TEntry", fieldbackground="#FFFFFF", bordercolor=COLOR_BORDER, lightcolor=COLOR_BORDER, darkcolor=COLOR_BORDER, padding=4)
        style.configure("TCombobox", fieldbackground="#FFFFFF", bordercolor=COLOR_BORDER, lightcolor=COLOR_BORDER, darkcolor=COLOR_BORDER, padding=4)
        style.configure("Treeview", font=("Segoe UI", 9), rowheight=26, fieldbackground="#FFFFFF", background="#FFFFFF", foreground=COLOR_TEXT, bordercolor=COLOR_BORDER)
        style.configure("Treeview.Heading", font=("Segoe UI", 9, "bold"), background="#EDF2F7", foreground=COLOR_TEXT)
        style.configure("TNotebook", background=COLOR_APP_BG, borderwidth=0)
        style.configure("TNotebook.Tab", font=("Segoe UI", 10), padding=(18, 8), background="#EAF0F7", foreground=COLOR_TEXT)
        style.map("TNotebook.Tab", background=[("selected", "#FFFFFF")], foreground=[("selected", COLOR_PRIMARY)])

        style.configure("TButton", font=("Segoe UI", 9), padding=(8, 7), background="#FFFFFF", foreground=COLOR_TEXT, bordercolor=COLOR_BORDER)
        style.map("TButton", background=[("active", "#EDF5FF")])
        style.configure("Primary.TButton", font=("Segoe UI", 9, "bold"), padding=(8, 8), background=COLOR_PRIMARY, foreground="white", bordercolor=COLOR_PRIMARY)
        style.map("Primary.TButton", background=[("active", COLOR_PRIMARY_DARK)], foreground=[("active", "white")])
        style.configure("Success.TButton", font=("Segoe UI", 9, "bold"), padding=(8, 8), background=COLOR_SUCCESS, foreground="white", bordercolor=COLOR_SUCCESS)
        style.map("Success.TButton", background=[("active", "#0B7A3E")], foreground=[("active", "white")])
        style.configure("Danger.TButton", font=("Segoe UI", 9), padding=(8, 7), background="#FFF1F0", foreground=COLOR_DANGER, bordercolor="#FFB4AB")
        style.map("Danger.TButton", background=[("active", "#FFE1DE")], foreground=[("active", COLOR_DANGER)])
        style.configure("Accent.TButton", font=("Segoe UI", 9, "bold"))

    def _build_left_panel(self, parent):
        top_row = ttk.Frame(parent, style="App.TFrame")
        top_row.pack(fill="x", pady=6)

        file_frame = ttk.LabelFrame(top_row, text="1. Imagen", style="Section.TLabelframe")
        file_frame.grid(row=0, column=0, sticky="nsew", padx=(0, 6))
        ttk.Button(file_frame, text="📂  Cargar imagen", command=self.load_image, style="Primary.TButton").pack(fill="x", padx=8, pady=8)
        self.image_label = ttk.Label(file_frame, text="No se ha cargado imagen.", wraplength=230)
        self.image_label.pack(fill="x", padx=8, pady=(0, 8))

        help_frame = ttk.LabelFrame(top_row, text="5. Ayuda rápida", style="Section.TLabelframe")
        help_frame.grid(row=0, column=1, sticky="nsew", padx=(6, 0))
        help_text = (
            "Use una sola calibración para ambas curvas.\n\n"
            "1) Calibre Xmin, Xmax, Ymin y Ymax.\n"
            "2) Seleccione los puntos de Mínima de Fusión.\n"
            "3) Seleccione los puntos de Total de Aclaramiento.\n"
            "4) Genere el Excel.\n\n"
            "Atajo: Ctrl+Z elimina el último punto activo."
        )
        ttk.Label(help_frame, text=help_text, wraplength=230, justify="left").pack(anchor="nw", padx=8, pady=8)

        top_row.columnconfigure(0, weight=1, uniform="top_cols")
        top_row.columnconfigure(1, weight=1, uniform="top_cols")

        meta_actions_row = ttk.Frame(parent, style="Panel.TFrame")
        meta_actions_row.pack(fill="x", pady=6)

        meta_frame = ttk.LabelFrame(meta_actions_row, text="2. Información del fusible", style="Section.TLabelframe")
        meta_frame.grid(row=0, column=0, sticky="nsew", padx=(0, 6))
        self._labeled_entry(meta_frame, "Nombre fusible", self.fuse_name)
        self._labeled_entry(meta_frame, "Hoja Excel", self.sheet_name)
        self._labeled_entry(meta_frame, "Etiqueta eje X", self.x_label)
        self._labeled_entry(meta_frame, "Unidad eje X", self.x_unit)
        self._labeled_entry(meta_frame, "Etiqueta eje Y", self.y_label)
        self._labeled_entry(meta_frame, "Unidad eje Y", self.y_unit)
        ttk.Checkbutton(meta_frame, text="Agregar hoja de configuración", variable=self.include_config_sheet).pack(anchor="w", padx=8, pady=(4, 8))

        actions_frame = ttk.LabelFrame(meta_actions_row, text="6. Acciones", style="Section.TLabelframe")
        actions_frame.grid(row=0, column=1, sticky="nsew", padx=(6, 0))
        ttk.Button(actions_frame, text="Eliminar último punto", command=self.delete_last_point, style="Danger.TButton").pack(fill="x", padx=8, pady=(8, 4))
        ttk.Button(actions_frame, text="Limpiar puntos activos", command=self.clear_active_curve_points).pack(fill="x", padx=8, pady=4)
        ttk.Button(actions_frame, text="Limpiar todos los puntos", command=self.clear_all_curve_points, style="Danger.TButton").pack(fill="x", padx=8, pady=4)
        ttk.Button(actions_frame, text="Limpiar calibración", command=self.clear_calibration).pack(fill="x", padx=8, pady=4)
        ttk.Button(actions_frame, text="📊  Generar Excel", command=self.export_to_excel, style="Success.TButton").pack(fill="x", padx=8, pady=(6, 8))

        meta_actions_row.columnconfigure(0, weight=1, uniform="meta_cols")
        meta_actions_row.columnconfigure(1, weight=1, uniform="meta_cols")

        axis_frame = ttk.LabelFrame(parent, text="3. Límites y escala", style="Section.TLabelframe")
        axis_frame.pack(fill="x", pady=6)

        ttk.Label(axis_frame, text="Escala X").grid(row=0, column=0, padx=8, pady=4, sticky="w")
        ttk.Combobox(axis_frame, textvariable=self.x_scale_type, values=["Logarítmica", "Lineal"], state="readonly", width=16).grid(row=0, column=1, padx=8, pady=4, sticky="ew")
        ttk.Label(axis_frame, text="Escala Y").grid(row=1, column=0, padx=8, pady=4, sticky="w")
        ttk.Combobox(axis_frame, textvariable=self.y_scale_type, values=["Logarítmica", "Lineal"], state="readonly", width=16).grid(row=1, column=1, padx=8, pady=4, sticky="ew")

        self._grid_entry(axis_frame, 2, "Xmin", self.x_min_entry)
        self._grid_entry(axis_frame, 3, "Xmax", self.x_max_entry)
        self._grid_entry(axis_frame, 4, "Ymin", self.y_min_entry)
        self._grid_entry(axis_frame, 5, "Ymax", self.y_max_entry)
        axis_frame.columnconfigure(1, weight=1)

        cal_frame = ttk.LabelFrame(parent, text="4. Calibración", style="Section.TLabelframe")
        cal_frame.pack(fill="x", pady=6)

        ttk.Label(cal_frame, text="Seleccione los puntos de referencia sobre la imagen.", wraplength=500).pack(anchor="w", padx=8, pady=(8, 4))
        ttk.Radiobutton(cal_frame, text="Marcar Xmin", variable=self.mode, value="x_min").pack(anchor="w", padx=8)
        ttk.Radiobutton(cal_frame, text="Marcar Xmax", variable=self.mode, value="x_max").pack(anchor="w", padx=8)
        ttk.Radiobutton(cal_frame, text="Marcar Ymin", variable=self.mode, value="y_min").pack(anchor="w", padx=8)
        ttk.Radiobutton(cal_frame, text="Marcar Ymax", variable=self.mode, value="y_max").pack(anchor="w", padx=8)
        ttk.Radiobutton(cal_frame, text="Seleccionar puntos de curva", variable=self.mode, value="seleccionar_puntos").pack(anchor="w", padx=8, pady=(0, 8))
        ttk.Button(cal_frame, text="Aplicar valores de ejes", command=self.apply_axis_values, style="Primary.TButton").pack(fill="x", padx=8, pady=(0, 8))

    def _labeled_entry(self, parent, label, variable):
        frame = ttk.Frame(parent, style="Panel.TFrame")
        frame.pack(fill="x", padx=8, pady=3)
        ttk.Label(frame, text=label, width=16).pack(side="left")
        ttk.Entry(frame, textvariable=variable).pack(side="left", fill="x", expand=True)

    def _grid_entry(self, parent, row, label, variable):
        ttk.Label(parent, text=label).grid(row=row, column=0, padx=8, pady=4, sticky="w")
        ttk.Entry(parent, textvariable=variable).grid(row=row, column=1, padx=8, pady=4, sticky="ew")

    def _build_canvas_area(self, parent):
        canvas_frame = ttk.LabelFrame(parent, text="Imagen y selección de puntos", style="Section.TLabelframe")
        canvas_frame.pack(fill="both", expand=True)

        toolbar = ttk.Frame(canvas_frame, style="Panel.TFrame")
        toolbar.pack(fill="x", padx=6, pady=6)
        ttk.Button(toolbar, text="Ajustar imagen", command=self.fit_image_to_canvas, style="Primary.TButton").pack(side="left", padx=3)
        ttk.Button(toolbar, text="Zoom +", command=lambda: self.zoom_image(1.15)).pack(side="left", padx=3)
        ttk.Button(toolbar, text="Zoom -", command=lambda: self.zoom_image(1 / 1.15)).pack(side="left", padx=3)
        self.active_curve_badge = ttk.Label(toolbar, text="Pestaña activa: Mínima de Fusión", font=("Segoe UI", 9, "bold"))
        self.active_curve_badge.pack(side="left", padx=12)
        self.status_label = ttk.Label(toolbar, text="Cargue una imagen para iniciar.")
        self.status_label.pack(side="left", padx=12)

        self.canvas = tk.Canvas(canvas_frame, bg="#FFFFFF", highlightthickness=1, highlightbackground=COLOR_BORDER)
        self.canvas.pack(fill="both", expand=True, padx=6, pady=(0, 6))
        self.canvas.bind("<Button-1>", self.on_canvas_click)
        self.canvas.bind("<Motion>", self.on_canvas_motion)
        self.canvas.bind("<Configure>", lambda event: self.redraw_canvas())

    def _build_tabs_area(self, parent):
        tab_frame = ttk.LabelFrame(parent, text="Puntos digitalizados", style="Section.TLabelframe")
        tab_frame.pack(fill="x", pady=(8, 0))

        self.notebook = ttk.Notebook(tab_frame)
        self.notebook.pack(fill="x", expand=True, padx=6, pady=6)
        self.notebook.bind("<<NotebookTabChanged>>", self.on_tab_changed)

        for curve_key in CURVE_KEYS:
            frame = ttk.Frame(self.notebook, style="Panel.TFrame")
            self.notebook.add(frame, text=curve_key)
            tree = self._create_points_tree(frame)
            self.trees[curve_key] = tree

    def _create_points_tree(self, parent):
        columns = ("No", "Time", "Current", "Pixel X", "Pixel Y")
        tree = ttk.Treeview(parent, columns=columns, show="headings", height=8)
        for col in columns:
            tree.heading(col, text=col)
        tree.column("No", width=50, anchor="center")
        tree.column("Time", width=150, anchor="e")
        tree.column("Current", width=150, anchor="e")
        tree.column("Pixel X", width=100, anchor="e")
        tree.column("Pixel Y", width=100, anchor="e")
        tree.pack(side="left", fill="x", expand=True, padx=(6, 0), pady=6)

        scrollbar = ttk.Scrollbar(parent, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=scrollbar.set)
        scrollbar.pack(side="right", fill="y", padx=(0, 6), pady=6)
        return tree

    def _bind_shortcuts(self):
        self.root.bind("<Control-z>", lambda event: self.delete_last_point())
        self.root.bind("<Control-Z>", lambda event: self.delete_last_point())

    def on_tab_changed(self, event=None):
        tab_index = self.notebook.index(self.notebook.select())
        curve_key = CURVE_KEYS[tab_index]
        self.active_curve.set(curve_key)
        self.active_curve_badge.config(text="Pestaña activa: " + curve_key)
        self.redraw_canvas()

    # -------------------------------------------------------------------------
    # IMAGEN / CANVAS
    # -------------------------------------------------------------------------
    def load_image(self):
        path = filedialog.askopenfilename(
            title="Seleccione imagen de curva",
            filetypes=[
                ("Imágenes", "*.png *.jpg *.jpeg *.bmp *.tif *.tiff"),
                ("Todos los archivos", "*.*"),
            ],
        )
        if not path:
            return
        try:
            self.original_image = Image.open(path).convert("RGB")
            self.image_path = path
            self.image_label.config(text=os.path.basename(path))
            self.clear_all_curve_points(confirm=False)
            self.clear_calibration(confirm=False)
            self.root.after(150, self.fit_image_to_canvas)
            self.set_status("Imagen cargada. Defina los límites y calibre los ejes.")
        except Exception as ex:
            messagebox.showerror("Error", "No fue posible cargar la imagen:\n" + str(ex))

    def fit_image_to_canvas(self):
        if self.original_image is None:
            return
        canvas_w = max(100, self.canvas.winfo_width() - 20)
        canvas_h = max(100, self.canvas.winfo_height() - 20)
        img_w, img_h = self.original_image.size
        scale = min(canvas_w / img_w, canvas_h / img_h)
        scale = max(0.05, min(scale, 10.0))
        self.image_scale = scale
        self.image_offset_x = 10
        self.image_offset_y = 10
        self.redraw_canvas()

    def zoom_image(self, factor):
        if self.original_image is None:
            return
        self.image_scale = max(0.05, min(self.image_scale * factor, 20.0))
        self.redraw_canvas()

    def redraw_canvas(self):
        self.canvas.delete("all")
        self.marker_ids.clear()
        self.calibration_marker_ids.clear()

        if self.original_image is None:
            self.canvas.create_text(
                self.canvas.winfo_width() / 2,
                self.canvas.winfo_height() / 2,
                text="Cargue una imagen de la curva TCC / fusible",
                fill="#606060",
                font=("Segoe UI", 14),
            )
            return

        img_w, img_h = self.original_image.size
        disp_w = max(1, int(img_w * self.image_scale))
        disp_h = max(1, int(img_h * self.image_scale))
        resized = self.original_image.resize((disp_w, disp_h), Image.LANCZOS)
        self.tk_image = ImageTk.PhotoImage(resized)
        self.image_on_canvas = self.canvas.create_image(self.image_offset_x, self.image_offset_y, anchor="nw", image=self.tk_image)

        self.draw_calibration_markers()
        self.draw_curve_markers()

    def image_to_canvas_coords(self, x_img, y_img):
        x = self.image_offset_x + x_img * self.image_scale
        y = self.image_offset_y + y_img * self.image_scale
        return x, y

    def canvas_to_image_coords(self, x_canvas, y_canvas):
        x = (x_canvas - self.image_offset_x) / self.image_scale
        y = (y_canvas - self.image_offset_y) / self.image_scale
        return x, y

    def point_inside_image(self, x_img, y_img):
        if self.original_image is None:
            return False
        img_w, img_h = self.original_image.size
        return 0 <= x_img <= img_w and 0 <= y_img <= img_h

    # -------------------------------------------------------------------------
    # EVENTOS
    # -------------------------------------------------------------------------
    def on_canvas_click(self, event):
        if self.original_image is None:
            messagebox.showinfo("Imagen requerida", "Primero cargue una imagen.")
            return

        x_img, y_img = self.canvas_to_image_coords(event.x, event.y)
        if not self.point_inside_image(x_img, y_img):
            return

        current_mode = self.mode.get()
        if current_mode in ["x_min", "x_max", "y_min", "y_max"]:
            self.set_calibration_point(current_mode, x_img, y_img)
            return

        if current_mode == "seleccionar_puntos":
            self.add_curve_point(x_img, y_img)

    def on_canvas_motion(self, event):
        if self.original_image is None:
            return
        x_img, y_img = self.canvas_to_image_coords(event.x, event.y)
        if not self.point_inside_image(x_img, y_img):
            self.set_status("Fuera de imagen.")
            return
        if self.is_calibration_complete():
            try:
                current_value, time_value = self.convert_pixel_to_values(x_img, y_img)
                self.set_status("Pixel: ({:.1f}, {:.1f}) | Current: {} | Time: {}".format(x_img, y_img, format_number(current_value), format_number(time_value)))
            except Exception:
                self.set_status("Pixel: ({:.1f}, {:.1f})".format(x_img, y_img))
        else:
            self.set_status("Pixel: ({:.1f}, {:.1f}) | Calibración pendiente.".format(x_img, y_img))

    # -------------------------------------------------------------------------
    # CALIBRACION
    # -------------------------------------------------------------------------
    def apply_axis_values(self):
        try:
            x_min = parse_float(self.x_min_entry.get(), "Xmin")
            x_max = parse_float(self.x_max_entry.get(), "Xmax")
            y_min = parse_float(self.y_min_entry.get(), "Ymin")
            y_max = parse_float(self.y_max_entry.get(), "Ymax")

            if x_min == x_max:
                raise ValueError("Xmin y Xmax no pueden ser iguales.")
            if y_min == y_max:
                raise ValueError("Ymin y Ymax no pueden ser iguales.")

            if self.x_scale_type.get() == "Logarítmica":
                validate_positive(x_min, "Xmin")
                validate_positive(x_max, "Xmax")
            if self.y_scale_type.get() == "Logarítmica":
                validate_positive(y_min, "Ymin")
                validate_positive(y_max, "Ymax")

            self.calibration["x_min_value"] = x_min
            self.calibration["x_max_value"] = x_max
            self.calibration["y_min_value"] = y_min
            self.calibration["y_max_value"] = y_max

            self.refresh_all_points()
            self.set_status("Valores de ejes aplicados.")
            messagebox.showinfo("Ejes actualizados", "Los valores de ejes fueron aplicados correctamente.")
        except Exception as ex:
            messagebox.showerror("Error en límites", str(ex))

    def set_calibration_point(self, mode, x_img, y_img):
        mapping = {
            "x_min": "x_min_pixel",
            "x_max": "x_max_pixel",
            "y_min": "y_min_pixel",
            "y_max": "y_max_pixel",
        }
        self.calibration[mapping[mode]] = (x_img, y_img)
        self.apply_axis_values_silent()
        self.redraw_canvas()
        self.set_status("Punto de calibración registrado: " + mode)

    def apply_axis_values_silent(self):
        try:
            self.calibration["x_min_value"] = parse_float(self.x_min_entry.get(), "Xmin")
            self.calibration["x_max_value"] = parse_float(self.x_max_entry.get(), "Xmax")
            self.calibration["y_min_value"] = parse_float(self.y_min_entry.get(), "Ymin")
            self.calibration["y_max_value"] = parse_float(self.y_max_entry.get(), "Ymax")
        except Exception:
            pass

    def is_calibration_complete(self):
        required = [
            "x_min_pixel", "x_max_pixel", "y_min_pixel", "y_max_pixel",
            "x_min_value", "x_max_value", "y_min_value", "y_max_value",
        ]
        return all(self.calibration.get(k) is not None for k in required)

    def convert_pixel_to_values(self, x_img, y_img):
        if not self.is_calibration_complete():
            raise ValueError("La calibración no está completa.")

        x_min_pixel = self.calibration["x_min_pixel"][0]
        x_max_pixel = self.calibration["x_max_pixel"][0]
        y_min_pixel = self.calibration["y_min_pixel"][1]
        y_max_pixel = self.calibration["y_max_pixel"][1]

        current_value = pixel_to_value(
            x_img,
            x_min_pixel,
            x_max_pixel,
            self.calibration["x_min_value"],
            self.calibration["x_max_value"],
            self.x_scale_type.get(),
        )
        time_value = pixel_to_value(
            y_img,
            y_min_pixel,
            y_max_pixel,
            self.calibration["y_min_value"],
            self.calibration["y_max_value"],
            self.y_scale_type.get(),
        )
        return current_value, time_value

    def draw_calibration_markers(self):
        labels = {
            "x_min_pixel": "Xmin",
            "x_max_pixel": "Xmax",
            "y_min_pixel": "Ymin",
            "y_max_pixel": "Ymax",
        }
        for key, label in labels.items():
            pt = self.calibration.get(key)
            if pt is None:
                continue
            x, y = self.image_to_canvas_coords(pt[0], pt[1])
            r = 6
            self.canvas.create_oval(x - r, y - r, x + r, y + r, outline="#005bbb", width=2, fill="#ffffff")
            self.canvas.create_text(x + 22, y - 10, text=label, fill="#005bbb", font=("Segoe UI", 9, "bold"))

    def clear_calibration(self, confirm=True):
        if confirm:
            if not messagebox.askyesno("Confirmar", "¿Desea borrar la calibración de ejes?"):
                return
        for key in self.calibration:
            if key.endswith("_pixel"):
                self.calibration[key] = None
        self.redraw_canvas()
        self.set_status("Calibración eliminada.")

    # -------------------------------------------------------------------------
    # PUNTOS DE CURVA
    # -------------------------------------------------------------------------
    def get_active_curve_key(self):
        return self.active_curve.get()

    def add_curve_point(self, x_img, y_img):
        if not self.is_calibration_complete():
            messagebox.showwarning("Calibración requerida", "Complete la calibración de ejes antes de seleccionar puntos de curva.")
            return
        try:
            current_value, time_value = self.convert_pixel_to_values(x_img, y_img)
            curve_key = self.get_active_curve_key()
            point = {
                "No": self.point_counter[curve_key],
                "Current": current_value,
                "Time": time_value,
                "Pixel X": x_img,
                "Pixel Y": y_img,
                "Curva": curve_key,
            }
            self.points[curve_key].append(point)
            self.point_counter[curve_key] += 1
            self.refresh_table(curve_key)
            self.redraw_canvas()
            self.set_status("Punto agregado en {}: Time={}, Current={}".format(curve_key, format_number(time_value), format_number(current_value)))
        except Exception as ex:
            messagebox.showerror("Error", str(ex))

    def draw_curve_markers(self):
        # Mostrar ambas curvas sobre la imagen; la activa queda con línea continua y la no activa en línea punteada.
        active = self.get_active_curve_key()
        for curve_key in CURVE_KEYS:
            color_cfg = CURVE_COLORS[curve_key]
            pts = self.points[curve_key]
            for point in pts:
                x, y = self.image_to_canvas_coords(point["Pixel X"], point["Pixel Y"])
                r = 4 if curve_key == active else 3
                self.canvas.create_oval(x - r, y - r, x + r, y + r, outline=color_cfg["line"], width=2, fill=color_cfg["fill"])
                self.canvas.create_text(x + 12, y - 8, text=str(point["No"]), fill=color_cfg["line"], font=("Segoe UI", 8, "bold"))

            if len(pts) >= 2:
                coords = []
                for point in pts:
                    x, y = self.image_to_canvas_coords(point["Pixel X"], point["Pixel Y"])
                    coords.extend([x, y])
                dash = None if curve_key == active else (3, 2)
                self.canvas.create_line(*coords, fill=color_cfg["line"], width=2 if curve_key == active else 1, dash=dash)

    def renumber_curve(self, curve_key):
        for idx, point in enumerate(self.points[curve_key], start=1):
            point["No"] = idx
        self.point_counter[curve_key] = len(self.points[curve_key]) + 1

    def delete_last_point(self):
        curve_key = self.get_active_curve_key()
        if not self.points[curve_key]:
            return
        self.points[curve_key].pop()
        self.renumber_curve(curve_key)
        self.refresh_table(curve_key)
        self.redraw_canvas()
        self.set_status("Último punto eliminado de " + curve_key + ".")

    def clear_active_curve_points(self, confirm=True):
        curve_key = self.get_active_curve_key()
        if confirm and self.points[curve_key]:
            if not messagebox.askyesno("Confirmar", "¿Desea borrar todos los puntos de " + curve_key + "?"):
                return
        self.points[curve_key] = []
        self.point_counter[curve_key] = 1
        self.refresh_table(curve_key)
        self.redraw_canvas()
        self.set_status("Puntos eliminados de " + curve_key + ".")

    def clear_all_curve_points(self, confirm=True):
        total_points = sum(len(self.points[k]) for k in CURVE_KEYS)
        if confirm and total_points > 0:
            if not messagebox.askyesno("Confirmar", "¿Desea borrar todos los puntos de ambas curvas?"):
                return
        for curve_key in CURVE_KEYS:
            self.points[curve_key] = []
            self.point_counter[curve_key] = 1
            self.refresh_table(curve_key)
        self.redraw_canvas()
        self.set_status("Puntos de ambas curvas eliminados.")

    def refresh_all_points(self):
        if not self.is_calibration_complete():
            for curve_key in CURVE_KEYS:
                self.refresh_table(curve_key)
            return
        try:
            for curve_key in CURVE_KEYS:
                for point in self.points[curve_key]:
                    current_value, time_value = self.convert_pixel_to_values(point["Pixel X"], point["Pixel Y"])
                    point["Current"] = current_value
                    point["Time"] = time_value
                    point["Curva"] = curve_key
                self.refresh_table(curve_key)
            self.redraw_canvas()
        except Exception:
            pass

    def refresh_table(self, curve_key):
        tree = self.trees.get(curve_key)
        if tree is None:
            return
        for item in tree.get_children():
            tree.delete(item)
        for point in self.points[curve_key]:
            tree.insert(
                "",
                "end",
                values=(
                    point["No"],
                    format_number(point["Time"]),
                    format_number(point["Current"]),
                    f"{point['Pixel X']:.1f}",
                    f"{point['Pixel Y']:.1f}",
                ),
            )

    # -------------------------------------------------------------------------
    # EXPORTACION
    # -------------------------------------------------------------------------
    def export_to_excel(self):
        total_points = sum(len(self.points[k]) for k in CURVE_KEYS)
        if total_points == 0:
            messagebox.showwarning("Sin datos", "No hay puntos digitalizados para exportar.")
            return
        if not self.is_calibration_complete():
            messagebox.showwarning("Calibración incompleta", "Complete la calibración antes de exportar.")
            return

        default_name = "curva_fusible_digitalizada.xlsx"
        path = filedialog.asksaveasfilename(
            title="Guardar Excel",
            defaultextension=".xlsx",
            initialfile=default_name,
            filetypes=[("Excel", "*.xlsx")],
        )
        if not path:
            return

        try:
            self.refresh_all_points()
            self._write_excel_like_template(path)
            messagebox.showinfo("Exportación completa", "Archivo generado correctamente:\n" + path)
            self.set_status("Excel exportado: " + os.path.basename(path))
        except Exception as ex:
            messagebox.showerror("Error al exportar", str(ex))

    def _write_excel_like_template(self, path):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
        from openpyxl.utils import get_column_letter
        from openpyxl.chart import ScatterChart, Series, Reference

        wb = Workbook()
        ws = wb.active
        ws.title = safe_sheet_name(self.sheet_name.get())

        fuse_name = self.fuse_name.get().strip() or "Fusible"

        # Estructura similar al archivo ejemplo.
        ws["A2"] = "Mínima de Fusión"
        ws["C2"] = "Total de Aclaramiento"
        ws["A3"] = fuse_name
        ws["C3"] = fuse_name
        ws["A4"] = "Time"
        ws["B4"] = "Current"
        ws["C4"] = "Time"
        ws["D4"] = "Current"

        min_points = self._points_sorted_for_export(CURVE_MIN_FUSION)
        total_points = self._points_sorted_for_export(CURVE_TOTAL_CLEARING)

        start_row = 5
        max_rows = max(len(min_points), len(total_points))
        for idx in range(max_rows):
            row = start_row + idx
            if idx < len(min_points):
                ws.cell(row=row, column=1, value=float(min_points[idx]["Time"]))
                ws.cell(row=row, column=2, value=float(min_points[idx]["Current"]))
            if idx < len(total_points):
                ws.cell(row=row, column=3, value=float(total_points[idx]["Time"]))
                ws.cell(row=row, column=4, value=float(total_points[idx]["Current"]))

        self._format_curve_sheet(ws, max_rows)
        self._add_curve_chart(ws, len(min_points), len(total_points))

        if self.include_config_sheet.get():
            cfg = wb.create_sheet("Configuración")
            self._write_config_sheet(cfg)

        wb.save(path)

    def _points_sorted_for_export(self, curve_key):
        # Se conserva el orden de selección. Para TCC usualmente conviene seleccionar de izquierda a derecha o por recorrido de curva.
        return list(self.points[curve_key])

    def _format_curve_sheet(self, ws, max_rows):
        from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
        from openpyxl.utils import get_column_letter

        title_fill = PatternFill("solid", fgColor="D9EAF7")
        header_fill = PatternFill("solid", fgColor="1F4E78")
        header_font = Font(color="FFFFFF", bold=True)
        normal_font = Font(name="Calibri", size=11)
        bold_font = Font(name="Calibri", size=11, bold=True)
        thin = Side(border_style="thin", color="BFBFBF")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for row in ws.iter_rows(min_row=1, max_row=max(4 + max_rows, 4), min_col=1, max_col=4):
            for cell in row:
                cell.font = normal_font
                cell.border = border
                cell.alignment = Alignment(horizontal="center", vertical="center")

        for cell_ref in ["A2", "B2", "C2", "D2", "A3", "B3", "C3", "D3"]:
            ws[cell_ref].fill = title_fill
            ws[cell_ref].font = bold_font

        for cell_ref in ["A4", "B4", "C4", "D4"]:
            ws[cell_ref].fill = header_fill
            ws[cell_ref].font = header_font

        # Simular encabezados por bloque sin fusionar celdas, para mantener compatibilidad con ETAP.
        ws["B2"] = None
        ws["B3"] = None
        ws["D2"] = None
        ws["D3"] = None

        for col in range(1, 5):
            ws.column_dimensions[get_column_letter(col)].width = 15

        for row in range(5, 5 + max_rows):
            for col in range(1, 5):
                ws.cell(row=row, column=col).number_format = "0.000000"
                if col in [2, 4]:
                    ws.cell(row=row, column=col).number_format = "0.000"

        ws.freeze_panes = "A5"
        ws.sheet_view.showGridLines = True

    def _add_curve_chart(self, ws, min_count, total_count):
        # Gráfico auxiliar en Excel. Los datos principales quedan en A:D como el archivo ejemplo.
        try:
            from openpyxl.chart import ScatterChart, Series, Reference

            chart = ScatterChart()
            chart.title = "Curvas de Fusible"
            chart.style = 13
            chart.x_axis.title = "Corriente [A]"
            chart.y_axis.title = "Tiempo [s]"
            chart.legend.position = "b"
            chart.height = 12
            chart.width = 20

            # Configuración logarítmica para TCC.
            chart.x_axis.scaling.logBase = 10
            chart.y_axis.scaling.logBase = 10

            if min_count > 0:
                yvalues = Reference(ws, min_col=1, min_row=5, max_row=4 + min_count)
                xvalues = Reference(ws, min_col=2, min_row=5, max_row=4 + min_count)
                s = Series(yvalues, xvalues, title="Mínima de Fusión")
                chart.series.append(s)

            if total_count > 0:
                yvalues = Reference(ws, min_col=3, min_row=5, max_row=4 + total_count)
                xvalues = Reference(ws, min_col=4, min_row=5, max_row=4 + total_count)
                s = Series(yvalues, xvalues, title="Total de Aclaramiento")
                chart.series.append(s)

            ws.add_chart(chart, "F2")
        except Exception:
            # El gráfico es auxiliar; si falla no debe impedir la exportación de datos.
            pass

    def _write_config_sheet(self, ws):
        from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
        from openpyxl.utils import get_column_letter

        rows = [
            ["Parámetro", "Valor"],
            ["Fecha de exportación", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
            ["Imagen", self.image_path or ""],
            ["Nombre de fusible", self.fuse_name.get().strip()],
            ["Hoja de datos", safe_sheet_name(self.sheet_name.get())],
            ["Eje X", self.x_label.get().strip()],
            ["Unidad X", self.x_unit.get().strip()],
            ["Escala X", self.x_scale_type.get()],
            ["Xmin", self.calibration["x_min_value"]],
            ["Xmax", self.calibration["x_max_value"]],
            ["Eje Y", self.y_label.get().strip()],
            ["Unidad Y", self.y_unit.get().strip()],
            ["Escala Y", self.y_scale_type.get()],
            ["Ymin", self.calibration["y_min_value"]],
            ["Ymax", self.calibration["y_max_value"]],
            ["Pixel Xmin", str(self.calibration["x_min_pixel"])],
            ["Pixel Xmax", str(self.calibration["x_max_pixel"])],
            ["Pixel Ymin", str(self.calibration["y_min_pixel"])],
            ["Pixel Ymax", str(self.calibration["y_max_pixel"])],
            ["Puntos Mínima de Fusión", len(self.points[CURVE_MIN_FUSION])],
            ["Puntos Total de Aclaramiento", len(self.points[CURVE_TOTAL_CLEARING])],
        ]
        for r, row in enumerate(rows, start=1):
            for c, value in enumerate(row, start=1):
                ws.cell(row=r, column=c, value=value)

        header_fill = PatternFill("solid", fgColor="1F4E78")
        header_font = Font(color="FFFFFF", bold=True)
        thin = Side(border_style="thin", color="BFBFBF")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        for row in ws.iter_rows(min_row=1, max_row=len(rows), min_col=1, max_col=2):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="center")
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 70
        ws.freeze_panes = "A2"

    # -------------------------------------------------------------------------
    # ESTADO
    # -------------------------------------------------------------------------
    def set_status(self, text):
        self.status_label.config(text=text)


# =============================================================================
# MAIN
# =============================================================================

def main():
    root = tk.Tk()
    app = FuseCurveDigitizerApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
